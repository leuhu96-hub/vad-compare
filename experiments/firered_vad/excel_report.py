"""Báo cáo Excel cho FireRedVAD theo đúng bố cục sheet "KQ du lieu that" (mục A, B, C) của file
Tong_hop_danh_gia_VAD_test.xlsx, để so trực tiếp với YAMNet.

    # chỉ có audio + GT + model: tự chia dev/test (30 % dev, theo file, phân tầng category, seed 3)
    python excel_report.py --audio-root "D:/data/phase_1" --gt-root "D:/data/Groundtruth" \
        --model-dir models -o firered_report.xlsx

    # toàn bộ là test, ngưỡng cố định
    python excel_report.py --audio-root ... --gt-root ... --model-dir models --dev-ratio 0 --threshold 0.5

    # dùng đúng chia dev/test của file YAMNet và thêm mục so sánh D
    python excel_report.py --xlsx Tong_hop_danh_gia_VAD_test.xlsx --audio-root ... --gt-root ... --model-dir models

Giữ nguyên quy trình của file YAMNet:
  * chia dev / test: đọc sheet "DL theo file" của --xlsx, hoặc --split-csv, hoặc tự chia (--dev-ratio)
  * GT bin 0.5 s (đa số thời lượng), cờ biên ±1 bin quanh điểm chuyển GT
  * quét ngưỡng 0.01 … 0.95 (bước 0.01), chọn ngưỡng trên dev theo DCF-min hoặc F1-max
  * DCF = w·MR + (1−w)·FAR (w = 0.75), CI 95% bootstrap theo file (1000 lần, seed 0)

Ba cấu hình FireRed (tương ứng A0 … A4 của YAMNet):
  F0  prob 10 ms ≥ ngưỡng, không hậu xử lý
  F1  hậu xử lý gốc FireRed: làm mượt 5 frame, máy trạng thái 0.2 s / 0.2 s, lùi đầu đoạn, cắt đoạn > 20 s
  F2  prob 10 ms ≥ ngưỡng + hậu xử lý A4 của YAMNet: pad 100 / 120 ms → drop → merge < 500 ms

Quyết định 10 ms → bin 0.5 s: --rebin overlap (mặc định; bin = speech nếu phủ ≥ --min-overlap, mặc định 20 %)
hoặc --rebin midpoint (tâm bin nằm trong đoạn speech).

Prob của model được cache trong --cache (chạy lại chỉ mất vài giây). Excel sinh ra dùng công thức
SUMIFS trên sheet "DL quet nguong": đổi cấu hình (B7), tiêu chí (B8), trọng số DCF (B6) là tự tính lại.
"""
from __future__ import annotations

import argparse
import csv
import datetime
import hashlib
import os
import sys
from collections import defaultdict

import numpy as np

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from dataset_reader import AUDIO_EXTS, file_key, find_audio, gt_to_bins, read_ground_truth  # noqa: E402

FRAME = 0.01
THRESHOLDS = [round(k * 0.01, 2) for k in range(1, 96)]
CONFIGS = ["F0", "F1", "F2"]
REGIONS = ("all", "nb", "b")
SHEET = "KQ FireRed"


# =========================================================================== split / matching
def read_split_xlsx(path: str):
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "DL theo file" not in wb.sheetnames:
        raise SystemExit(f"{path}: không có sheet 'DL theo file' để lấy danh sách dev/test")
    ws = wb["DL theo file"]
    rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True))
    hdr = next((i for i, r in enumerate(rows) if r and "File" in r and "Tập" in r), None)
    if hdr is None:
        raise SystemExit(f"{path}: sheet 'DL theo file' không có header 'File' / 'Category' / 'Tập'")
    h = rows[hdr]
    ci, cc, cs = h.index("File"), h.index("Category"), h.index("Tập")
    out = []
    for r in rows[hdr + 1:]:
        if not r or not r[ci]:
            continue
        out.append({"file": str(r[ci]).strip(), "category": str(r[cc]).strip(), "split": str(r[cs]).strip()})
    return out


def read_split_csv(path: str):
    out = []
    with open(path, encoding="utf-8-sig") as f:
        for r in csv.DictReader(f):
            low = {k.strip().lower(): (v or "").strip() for k, v in r.items()}
            name = low.get("file") or low.get("name")
            if not name:
                continue
            out.append({"file": name, "category": low.get("category", ""),
                        "split": low.get("tập") or low.get("split") or low.get("set") or "test"})
    return out


def index_gt(gt_root: str):
    """stem -> [paths]; nếu một tên có nhiều file, ưu tiên file nằm trong thư mục tên ground_truth."""
    idx = defaultdict(list)
    for dp, _, files in os.walk(gt_root):
        for f in files:
            if f.lower().endswith(".txt"):
                idx[file_key(f)].append(os.path.join(dp, f))
    out = {}
    for k, paths in idx.items():
        pref = [p for p in paths if "ground_truth" in [x.lower() for x in os.path.abspath(p).split(os.sep)]]
        out[k] = pref or paths
    return out


def match(split, audio_root, gt_root, quality=None, log=print):
    audio = defaultdict(list)
    for p, q in find_audio(audio_root, AUDIO_EXTS):
        audio[file_key(p)].append(p)
    gts = index_gt(gt_root)
    items, problems = [], []
    for s in split:
        k = file_key(s["file"])
        a = audio.get(k, [])
        if quality and len(a) > 1:
            a = [p for p in a if quality.lower() in p.lower()] or a
        g = gts.get(k, [])
        if not a or not g:
            problems.append((s["file"], "không có audio" if not a else "không có GT"))
            continue
        if len(g) > 1:
            problems.append((s["file"], f"{len(g)} file GT trùng tên: {g[:2]}"))
            continue
        if len(a) > 1:
            log(f"[!] {s['file']}: {len(a)} file audio trùng tên, dùng {a[0]} (chọn bản khác bằng --quality)")
        items.append({**s, "audio": a[0], "gt": g[0]})
    return items, problems


def gt_category(gt_path, gt_root):
    """category = thư mục đầu tiên dưới gt_root (Groundtruth/<category>/ground_truth/x.txt);
    GT để phẳng thì lấy tiền tố tên file trước dấu '_' (Asm_xxx -> Asm)."""
    rel = os.path.relpath(gt_path, gt_root).split(os.sep)
    if len(rel) > 1 and rel[0].lower() != "ground_truth":
        return rel[0]
    stem = os.path.splitext(os.path.basename(gt_path))[0]
    return stem.split("_", 1)[0] if "_" in stem else "Tất cả"


def auto_split(audio_root, gt_root, dev_ratio, seed):
    """Không có file chia sẵn: mọi file có cả audio và GT; chia dev/test theo file, phân tầng theo category."""
    audio = {file_key(p) for p, _ in find_audio(audio_root, AUDIO_EXTS)}
    by_cat = defaultdict(list)
    for k, paths in sorted(index_gt(gt_root).items()):
        if k in audio:
            name = os.path.splitext(os.path.basename(paths[0]))[0]
            by_cat[gt_category(paths[0], gt_root)].append(name)
    rng = np.random.default_rng(seed)
    out = []
    for cat in sorted(by_cat):
        names = sorted(by_cat[cat])
        n_dev = int(round(dev_ratio * len(names)))
        if dev_ratio > 0 and len(names) >= 2:
            n_dev = min(max(n_dev, 1), len(names) - 1)
        dev = set(rng.permutation(len(names))[:n_dev].tolist())
        out += [{"file": n, "category": cat, "split": "dev" if i in dev else "test"} for i, n in enumerate(names)]
    return out


def write_split(path, split):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["File", "Category", "Tập"])
        for s in split:
            w.writerow([s["file"], s["category"], s["split"]])


# =========================================================================== FireRed prob (cache)
def get_probs(items, args, log=print):
    model = None
    os.makedirs(args.cache, exist_ok=True)
    for i, it in enumerate(items, 1):
        h = hashlib.md5(f"{os.path.abspath(it['audio'])}|{args.onnx or args.model_dir}|{args.mode}".encode()).hexdigest()[:8]
        cp = os.path.join(args.cache, f"{file_key(it['audio'])}_{h}.npz")
        if os.path.exists(cp) and not args.force:
            z = np.load(cp)
            it["probs"], it["dur"] = z["probs"].astype(np.float64), float(z["duration"])
            continue
        if model is None:
            from firered_onnx import FireRedOnnx

            model = FireRedOnnx(model_dir=args.model_dir, mode=args.mode, threads=args.threads,
                                onnx_path=args.onnx, cmvn_path=args.cmvn)
            log(f"[model] {model.onnx_path} ({model.mode})")
        r = model.run_file(it["audio"])
        it["probs"], it["dur"] = r.probs.astype(np.float64), float(r.duration)
        np.savez_compressed(cp, probs=r.probs.astype(np.float32), duration=r.duration)
        log(f"[{i}/{len(items)}] {os.path.basename(it['audio'])} ({r.duration:.1f} s) "
            f"onnx {r.timings.get('7_onnx', 0):.2f} s")


# =========================================================================== post-processing (10 ms)
def _runs(b):
    """-> starts, ends (exclusive) của các đoạn 1 trong mảng 0/1."""
    d = np.diff(np.r_[0, b.astype(np.int8), 0])
    return np.flatnonzero(d == 1), np.flatnonzero(d == -1)


def _from_runs(n, starts, ends):
    out = np.zeros(n, np.int8)
    for s, e in zip(starts, ends):
        out[s:e] = 1
    return out


def fsm(b, min_sp, min_sil):
    """Máy trạng thái FireRed (giống _state_machine của upstream), cài theo run-length."""
    n = len(b)
    if min_sp <= 0 and min_sil <= 0:
        return b.astype(np.int8)
    msp, msil = max(min_sp, 1), max(min_sil, 1)
    dec = np.zeros(n, np.int8)
    idx = np.flatnonzero(np.diff(b)) + 1
    speaking = False
    for s, e in zip(np.r_[0, idx], np.r_[idx, n]):
        L = e - s
        if not speaking:
            if b[s] == 1 and L > msp:
                dec[s:e] = 1
                speaking = True
        elif b[s] == 1:
            dec[s:e] = 1
        elif L > msil:
            dec[s:s + msil] = 1
            speaking = False
        else:
            dec[s:e] = 1
    return dec


def fix_smooth_start(dec, w):
    if w <= 1 or len(dec) < 2:
        return dec
    new = dec.copy()
    for t in np.flatnonzero((dec[1:] == 1) & (dec[:-1] == 0)) + 1:
        new[max(0, t - w):t] = 1
    return new


def split_long(dec, probs, max_frame):
    """Giống _split_long của upstream (kể cả cách đổi giây -> frame bằng int())."""
    starts, ends = _runs(dec)
    if not len(starts):
        return dec
    T = len(dec)
    new = dec.copy()
    for j, (s0, e0) in enumerate(zip(starts, ends)):
        a = int(round(s0 * FRAME, 3) / FRAME)
        e_s = round(T * FRAME + 0.025, 3) if e0 == T else round(e0 * FRAME, 3)
        b = int(e_s / FRAME)
        if b - a <= max_frame:
            continue
        seg, start = probs[a:b], 0
        while len(seg) - start > max_frame:
            w0, w1 = int(start + max_frame / 2), int(start + max_frame)
            k = w0 + int(np.argmin(seg[w0:w1]))
            new[a + k] = 0
            start = k + 1
    return new


def smooth(p, w):
    from firered_onnx import _smooth

    return _smooth(p, w)


def firered_post(sm, probs, thr, post):
    b = (sm >= thr).astype(np.int8)
    dec = fsm(b, post.min_speech_frame, post.min_silence_frame)
    dec = fix_smooth_start(dec, max(1, post.smooth_window_size))
    return split_long(dec, probs, post.max_speech_frame)


def rules_post(on, pad_b, pad_a, drop, merge, order):
    """Hậu xử lý kiểu YAMNet A4 trên frame 10 ms (đơn vị frame). merge: lấp khoảng hở < merge."""
    T = len(on)
    s, e = _runs(on)
    for step in order:
        if not len(s):
            break
        if step == "pad":
            s, e = np.maximum(s - pad_b, 0), np.minimum(e + pad_a, T)
            keep_s, keep_e = [s[0]], [e[0]]
            for a, b in zip(s[1:], e[1:]):
                if a <= keep_e[-1]:
                    keep_e[-1] = max(keep_e[-1], b)
                else:
                    keep_s.append(a)
                    keep_e.append(b)
            s, e = np.array(keep_s), np.array(keep_e)
        elif step == "drop":
            k = (e - s) >= drop
            s, e = s[k], e[k]
        elif step == "merge":
            keep_s, keep_e = [s[0]], [e[0]]
            for a, b in zip(s[1:], e[1:]):
                if a - keep_e[-1] < merge:
                    keep_e[-1] = max(keep_e[-1], b)
                else:
                    keep_s.append(a)
                    keep_e.append(b)
            s, e = np.array(keep_s), np.array(keep_e)
    return _from_runs(T, s, e)


# =========================================================================== bins / counts
def boundary_mask(labels, tol=1):
    m = np.zeros(len(labels), bool)
    if tol <= 0:
        return m
    for t in range(1, len(labels)):
        a, b = labels[t - 1], labels[t]
        if a >= 0 and b >= 0 and a != b:
            m[max(0, t - tol):min(len(labels), t + tol)] = True
    return m


class Binner:
    """Quyết định 10 ms -> bin GT.
    Frame cuối phủ thêm 25 ms (cửa sổ 25 ms), giống decision_to_segment của FireRed: đoạn speech chạm
    frame cuối kéo tới min(T·10 ms + 25 ms, độ dài file). Nhờ vậy khớp từng bin với run.py."""

    def __init__(self, n, dur, bin_s, rule, min_ov, n_frames):
        self.n, self.fpb, self.T = n, int(round(bin_s / FRAME)), n_frames
        edges = np.arange(n + 1) * bin_s
        self.width = np.minimum(edges[1:], dur) - edges[:-1]
        self.rule, self.min_ov = rule, min_ov
        mid = (edges[:-1] + np.minimum(edges[1:], dur)) / 2
        self.mid = np.floor(mid / FRAME + 1e-9).astype(int)
        self.mid_t = mid
        t0 = round(n_frames * FRAME, 3)                       # giống decisions_to_segments (làm tròn ms)
        self.tail_end = min(round(t0 + 0.025, 3), dur)
        self.tail_start = t0
        self.tail_cov = np.clip(np.minimum(edges[1:], self.tail_end) - np.maximum(edges[:-1], t0), 0, None)

    def _pad(self, x, fill=0):
        L = self.n * self.fpb
        out = np.full(L, fill, dtype=x.dtype)
        m = min(L, len(x))
        out[:m] = x[:m]
        return out

    def _mid_vals(self, x, fill):
        """Giá trị tại tâm bin; tâm rơi vào đuôi 25 ms sau frame cuối thì lấy frame cuối."""
        v = np.full(self.n, fill, dtype=np.float64)
        inside = self.mid < len(x)
        v[inside] = x[self.mid[inside]]
        tail = ~inside & (self.mid_t >= self.tail_start) & (self.mid_t < self.tail_end) & (len(x) > 0)
        if tail.any():
            v[tail] = x[-1]
        return v

    def bins(self, dec):
        if self.rule == "midpoint":
            return self._mid_vals(dec, 0) > 0
        cov = self._pad(dec).reshape(self.n, self.fpb).sum(1) * FRAME
        if len(dec) and dec[-1]:
            cov = cov + self.tail_cov
        return cov >= self.min_ov * self.width - 1e-9

    def bins_all_thr(self, probs, thr):
        """F0 cho mọi ngưỡng một lần: -> [n_thr, n] bool."""
        if self.rule == "midpoint":
            return self._mid_vals(probs, -1.0)[None, :] >= thr[:, None]
        P = self._pad(probs, -1.0).reshape(self.n, self.fpb)
        cov = (P[None, :, :] >= thr[:, None, None]).sum(2) * FRAME
        if len(probs):
            cov = cov + (probs[-1] >= thr)[:, None] * self.tail_cov[None, :]
        return cov >= self.min_ov * self.width[None, :] - 1e-9


def counts(D, labels, bnd):
    """D [n_thr, n] bool -> [n_thr, 2 (b, nb), 4 (TP, FN, FP, TN)]."""
    y = labels == 1
    valid = labels >= 0
    out = np.zeros((D.shape[0], 2, 4), np.int64)
    for r, m in enumerate((valid & bnd, valid & ~bnd)):
        sp, ns = y & m, ~y & m
        out[:, r, 0] = (D & sp).sum(1)
        out[:, r, 1] = (~D & sp).sum(1)
        out[:, r, 2] = (D & ns).sum(1)
        out[:, r, 3] = (~D & ns).sum(1)
    return out


def evaluate_file(it, args, post):
    gt = read_ground_truth(it["gt"])
    labels = np.array(gt_to_bins(gt, args.bin, it["dur"]), dtype=int)
    n = len(labels)
    it["labels"], it["bnd"] = labels, boundary_mask(labels, args.boundary)
    if n == 0:
        it["counts"] = {c: np.zeros((len(THRESHOLDS), 2, 4), np.int64) for c in CONFIGS}
        return
    B = Binner(n, it["dur"], args.bin, args.rebin, args.min_overlap, len(it["probs"]))
    thr = np.array(THRESHOLDS)
    p = it["probs"]
    res = {"F0": counts(B.bins_all_thr(p, thr), labels, it["bnd"])}
    sm = smooth(p, max(1, post.smooth_window_size))
    D1 = np.stack([B.bins(firered_post(sm, p, t, post)) for t in THRESHOLDS])
    res["F1"] = counts(D1, labels, it["bnd"])
    fr = lambda s: int(round(s / FRAME))  # noqa: E731
    D2 = np.stack([B.bins(rules_post((p >= t).astype(np.int8), fr(args.f2_pad_before), fr(args.f2_pad_after),
                                     fr(args.f2_drop), fr(args.f2_merge), args.f2_order)) for t in THRESHOLDS])
    res["F2"] = counts(D2, labels, it["bnd"])
    it["counts"] = res


# =========================================================================== metrics / bootstrap
def metr(tp, fn, fp, tn, w):
    tp, fn, fp, tn = (np.asarray(x, float) for x in (tp, fn, fp, tn))
    with np.errstate(divide="ignore", invalid="ignore"):
        mr = np.where(tp + fn > 0, fn / (tp + fn), np.nan)
        far = np.where(fp + tn > 0, fp / (fp + tn), np.nan)
        f1 = np.where(2 * tp + fn + fp > 0, 2 * tp / (2 * tp + fn + fp), np.nan)
    return {"mr": mr, "far": far, "f1": f1, "dcf": w * mr + (1 - w) * far}


def choose_threshold(items, cfg, crit, w):
    dev = [it for it in items if it["split"] == "dev"]
    if not dev:
        return None, None
    c = sum(it["counts"][cfg] for it in dev).sum(1)        # [n_thr, 4] (all)
    m = metr(c[:, 0], c[:, 1], c[:, 2], c[:, 3], w)
    k = int(np.nanargmax(m["f1"])) if crit == "F1-max" else int(np.nanargmin(m["dcf"]))
    return k, THRESHOLDS[k]


def bootstrap(per_file, w, n_boot, rng):
    """per_file [nf, 2, 4] -> dict metric -> (lo, hi) cho all và nb."""
    nf = len(per_file)
    if nf == 0:
        return {}
    idx = rng.integers(0, nf, (n_boot, nf))
    s = per_file[idx].sum(1)                               # [B, 2, 4]
    a = s.sum(1)
    m = metr(a[:, 0], a[:, 1], a[:, 2], a[:, 3], w)
    nb = s[:, 1]
    m_nb = metr(nb[:, 0], nb[:, 1], nb[:, 2], nb[:, 3], w)
    ci = lambda x: (float(np.nanpercentile(x, 2.5)), float(np.nanpercentile(x, 97.5)))  # noqa: E731
    return {"f1": ci(m["f1"]), "dcf": ci(m["dcf"]), "mr": ci(m["mr"]), "far": ci(m["far"]),
            "f1_nb": ci(m_nb["f1"])}


# =========================================================================== YAMNet reference
def read_yamnet(path):
    """Số đếm test của YAMNet từ sheet 'KQ du lieu that' (giá trị đã lưu). None nếu không đọc được."""
    import openpyxl

    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        ws = wb["KQ du lieu that"]
    except Exception:
        return None
    rows = list(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 80), values_only=True))
    out = {"cfg": None, "thr": None, "strict": None, "cats": {}, "nfile": {}}
    for r in rows:
        if r and r[0] == "Cấu hình hậu xử lý cho tầng B, category":
            out["cfg"] = r[1]
        if r and r[0] == "Ngưỡng dùng (tự chọn trên dev theo B7, B8)":
            out["thr"] = r[1]
        if r and r[0] == "Strict (mọi bin)" and all(isinstance(x, (int, float)) for x in r[2:6]):
            out["strict"] = [int(x) for x in r[2:6]]
    try:
        hc = next(i for i, r in enumerate(rows) if r and r[0] == "Category" and r[2] == "TP")
    except StopIteration:
        return out if out["strict"] else None
    for r in rows[hc + 1:]:
        if not r or not r[0] or str(r[0]).startswith(("Tổng", "Macro")):
            break
        if all(isinstance(x, (int, float)) for x in r[2:6]):
            out["cats"][str(r[0])] = [int(x) for x in r[2:6]]
            out["nfile"][str(r[0])] = r[1]
    return out if out["strict"] else None


# =========================================================================== Excel
def write_xlsx(path, items, problems, args, sel, ci, ci_cat, yam):
    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    NAVY, BLUE, YEL = "1F4E78", "0000FF", "FFF2CC"
    F = lambda **k: Font(name="Arial", size=k.pop("size", 10), **k)  # noqa: E731
    title, sect = F(size=14, bold=True, color=NAVY), F(size=11, bold=True, color=NAVY)
    hdr_font, hdr_fill = F(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=NAVY)
    yel = PatternFill("solid", fgColor=YEL)
    blue, note = F(color=BLUE), F(size=9, color="595959")
    PCT, R3, R2 = "0.0%", "0.000", "0.00"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET
    dq = wb.create_sheet("DL quet nguong")
    df = wb.create_sheet("DL theo file")
    Q = "'DL quet nguong'"

    def put(sheet, ref, v, font=None, fmt=None, fill=None, wrap=False):
        c = sheet[ref]
        c.value = v
        c.font = font or F()
        if fmt:
            c.number_format = fmt
        if fill:
            c.fill = fill
        if wrap:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        return c

    def header(sheet, row, names, col=1):
        for j, h in enumerate(names):
            c = sheet.cell(row, col + j, h)
            c.font, c.fill = hdr_font, hdr_fill
            c.alignment = Alignment(wrap_text=True, vertical="center")

    # ------------------------------------------------------------- DL quet nguong
    put(dq, "A1", "Số đếm theo ngưỡng – FireRedVAD", sect)
    put(dq, "A2", "Tập: dev / test (theo sheet DL theo file). Vùng bin: all = mọi bin, nb = bỏ bin biên ±1, "
                  "b = chỉ bin biên. Chữ xanh = số đếm từ Python; MR…DCF (chỉ dev) là công thức.", note)
    cols = ["Tập", "Cấu hình", "Category", "Vùng bin", "Ngưỡng", "TP", "FN", "FP", "TN", "FEC", "MSC", "OVER",
            "NDS", "MR", "FAR", "Precision", "F1", "DCF"]
    header(dq, 4, cols)
    row = 5
    dev_rows = {}
    cats = sorted({it["category"] for it in items})

    def agg(sel_items, cfg):
        if not sel_items:
            return np.zeros((len(THRESHOLDS), 2, 4), np.int64)
        return sum(it["counts"][cfg] for it in sel_items)

    def write_row(split, cfg, cat, region, k, c):
        nonlocal row
        b, nb = c[k, 0], c[k, 1]
        if region == "all":
            v = b + nb
            extra = [b[1], nb[1], b[2], nb[2]]
        elif region == "nb":
            v, extra = nb, [0, nb[1], 0, nb[2]]
        else:
            v, extra = b, [b[1], 0, b[2], 0]
        vals = [split, cfg, cat, region, THRESHOLDS[k]] + [int(x) for x in v] + [int(x) for x in extra]
        for j, x in enumerate(vals):
            cell = dq.cell(row, j + 1, x)
            cell.font = blue if j >= 5 else F()
        if split == "dev":
            r = row
            dq.cell(r, 14, f'=IF(F{r}+G{r}=0,"",G{r}/(F{r}+G{r}))')
            dq.cell(r, 15, f'=IF(H{r}+I{r}=0,"",H{r}/(H{r}+I{r}))')
            dq.cell(r, 16, f'=IF(F{r}+H{r}=0,"",F{r}/(F{r}+H{r}))')
            dq.cell(r, 17, f'=IF(2*F{r}+G{r}+H{r}=0,"",2*F{r}/(2*F{r}+G{r}+H{r}))')
            dq.cell(r, 18, f"=IF(OR(N{r}=\"\",O{r}=\"\"),\"\",'{SHEET}'!$B$6*N{r}+(1-'{SHEET}'!$B$6)*O{r})")
            for j, fm in zip(range(14, 19), (PCT, PCT, PCT, R3, R3)):
                dq.cell(r, j).font = F()
                dq.cell(r, j).number_format = fm
        row += 1

    dev = [it for it in items if it["split"] == "dev"]
    test = [it for it in items if it["split"] == "test"]
    for cfg in CONFIGS:
        c = agg(dev, cfg)
        dev_rows[cfg] = (row, row + len(THRESHOLDS) - 1)
        for k in range(len(THRESHOLDS)):
            write_row("dev", cfg, "Tất cả", "all", k, c)
    for cfg in CONFIGS:
        for cat in cats:
            c = agg([it for it in test if it["category"] == cat], cfg)
            for region in REGIONS:
                for k in range(len(THRESHOLDS)):
                    write_row("test", cfg, cat, region, k, c)
    last = row - 1
    dq.freeze_panes = "A5"
    for j, wdt in enumerate([7, 9, 10, 9, 8] + [8] * 13, 1):
        dq.column_dimensions[openpyxl.utils.get_column_letter(j)].width = wdt

    def rng(col):
        return f"{Q}!${col}$5:${col}${last}"

    def sumifs(col, region_ref, thr_ref, cat_ref=None):
        s = (f'=SUMIFS({rng(col)},{rng("A")},"test",{rng("B")},$B$7,{rng("D")},{region_ref},'
             f'{rng("E")},{thr_ref}')
        return s + (f",{rng('C')},{cat_ref})" if cat_ref else ")")

    # ------------------------------------------------------------- DL theo file
    put(df, "A1", f"Số liệu theo file ({len(items)} file)", sect)
    put(df, "A2", f"Chữ xanh = số từ excel_report.py. Cột I–O tính ở cấu hình {sel['cfg']} và ngưỡng "
                  f"{sel['thr']} (lựa chọn của Python, xem '{SHEET}'!B12).", note)
    header(df, 4, ["File", "Category", "Tập", "Thời lượng (s)", "% speech (GT)", "#bin có nhãn", "#chuyển trạng thái",
                   "% bin biên", "TP", "FN", "FP", "TN", "MR", "FAR", "F1", "Audio", "GT"])
    for i, it in enumerate(items, 5):
        lab = it["labels"]
        valid = lab >= 0
        nlab = int(valid.sum())
        trans = int(np.sum((lab[1:] != lab[:-1]) & (lab[1:] >= 0) & (lab[:-1] >= 0))) if len(lab) > 1 else 0
        c = it["counts"][sel["cfg"]][sel["k"]].sum(0)
        m = metr(*c, args.dcf_weight)
        vals = [it["file"], it["category"], it["split"], round(it["dur"], 3),
                float((lab == 1).sum() / nlab) if nlab else None, nlab, trans,
                float((it["bnd"] & valid).sum() / nlab) if nlab else None,
                *[int(x) for x in c], *[None if np.isnan(m[x]) else float(m[x]) for x in ("mr", "far", "f1")],
                it["audio"], it["gt"]]
        for j, v in enumerate(vals, 1):
            cell = df.cell(i, j, v)
            cell.font = blue if 4 <= j <= 15 else F()
            if j in (5, 8, 13, 14):
                cell.number_format = PCT
            elif j == 15:
                cell.number_format = R3
    df_last = 4 + len(items)
    for j, wdt in enumerate([24, 10, 7, 11, 11, 9, 10, 9, 7, 7, 7, 7, 8, 8, 8, 50, 50], 1):
        df.column_dimensions[openpyxl.utils.get_column_letter(j)].width = wdt
    df.freeze_panes = "A5"

    # ------------------------------------------------------------- main sheet
    put(ws, "A1", "Kết quả đánh giá FireRedVAD – chỉ số tính bằng công thức", title)
    put(ws, "A2", f"Nguồn: excel_report.py ({datetime.date.today():%d/%m/%Y}). {len(items)} file "
                  f"(dev {len(dev)} / test {len(test)}); GT: {args.gt_root}; model: {args.onnx or args.model_dir}. "
                  f"Chia dev/test lấy từ {args.split_source}.", note)
    put(ws, "A3", "Nền vàng = tham số bạn chỉnh (B6:B8, B10, B14) · chữ đen = công thức · chữ xanh = số nhập từ Python "
                  "(bootstrap; không tính lại trong Excel). Đổi tham số thì bảng A, B, C, D tự tính lại.", note)
    rule = (f"phủ ≥ {args.min_overlap:.0%} bin" if args.rebin == "overlap" else "tâm bin nằm trong đoạn speech")
    put(ws, "A4", f"Quyết định 10 ms → bin {args.bin:g} s: {rule}. Biên = ±{args.boundary} bin quanh điểm chuyển GT. "
                  f"F2 = pad {args.f2_pad_before * 1000:.0f}/{args.f2_pad_after * 1000:.0f} ms, drop "
                  f"{args.f2_drop * 1000:.0f} ms, merge < {args.f2_merge * 1000:.0f} ms, thứ tự "
                  f"{' → '.join(args.f2_order)}.", note)
    put(ws, "A5", "Tham số", sect)
    params = [
        (6, "Trọng số miss trong DCF (FA = 1 − ô này)", args.dcf_weight, R2, True),
        (7, "Cấu hình hậu xử lý cho tầng B, category", sel["cfg"], None, True),
        (8, "Tiêu chí chọn ngưỡng trên dev", sel["crit"], None, True),
        (9, "Quy tắc rebin (cố định khi chạy script)", rule, None, False),
        (10, "Ngưỡng cố định (--threshold; mặc định 0.4 của FireRed)", args.fixed_thr, "0.00##", True),
    ]
    for r, lab, v, fmt, edit in params:
        put(ws, f"A{r}", lab, wrap=True)
        put(ws, f"B{r}", v, fmt=fmt, fill=yel if edit else None)
    dv1 = DataValidation(type="list", formula1='"F0,F1,F2"', allow_blank=False)
    dv2 = DataValidation(type="list", formula1='"DCF-min,F1-max"', allow_blank=False)
    ws.add_data_validation(dv1)
    ws.add_data_validation(dv2)
    dv1.add("B7")
    dv2.add("B8")
    put(ws, "A11", "Ngưỡng dùng (B14 = dev: chọn trên dev theo B7, B8; B14 = cố định: B10)", wrap=True)
    put(ws, "B11", '=IF($B$14="cố định",$B$10,INDEX($G$17:$G$19,MATCH($B$7,$A$17:$A$19,0)))' if dev else "=$B$10",
        F(bold=True), "0.00##")
    put(ws, "A14", "Cách lấy ngưỡng (dev | cố định)", wrap=True)
    put(ws, "B14", sel["mode"], fill=yel)
    dv3 = DataValidation(type="list", formula1='"dev,cố định"', allow_blank=False)
    ws.add_data_validation(dv3)
    dv3.add("B14")
    put(ws, "A12", "Ngưỡng Python đã dùng cho CI và sheet DL theo file", wrap=True)
    put(ws, "B12", sel["thr"], blue, "0.00##")
    put(ws, "A13", "Kiểm tra")
    put(ws, "B13", f'=IF(AND(ABS(B11-B12)<0.0000001,B7="{sel["cfg"]}",ABS(B6-{args.dcf_weight})<0.0000001),"Khớp",'
                   f'"KHÁC: CI và DL theo file vẫn ở {sel["cfg"]} / ngưỡng {sel["thr"]} / w {args.dcf_weight}")')

    put(ws, "A15", "A. Chọn ngưỡng trên dev (qua toàn pipeline; DCF theo trọng số B6)", sect)
    header(ws, 16, ["Cấu hình", "Mô tả", "Ngưỡng DCF-min (dev)", "DCF min (dev)", "Ngưỡng F1-max (dev)",
                    "F1 max (dev)", "Ngưỡng dùng (theo B8)"])
    desc = {
        "F0": "Prob 10 ms ≥ ngưỡng, không hậu xử lý",
        "F1": "Hậu xử lý gốc FireRed: làm mượt 5 frame, máy trạng thái 0.2 s / 0.2 s, lùi đầu đoạn 50 ms, "
              "cắt đoạn > 20 s",
        "F2": f"Prob 10 ms ≥ ngưỡng + hậu xử lý A4 của YAMNet (pad {args.f2_pad_before * 1000:.0f}/"
              f"{args.f2_pad_after * 1000:.0f} ms, merge < {args.f2_merge * 1000:.0f} ms)",
    }
    for r, cfg in zip((17, 18, 19), CONFIGS):
        a, b = dev_rows[cfg]
        E, Rr, Qq = f"{Q}!$E${a}:$E${b}", f"{Q}!$R${a}:$R${b}", f"{Q}!$Q${a}:$Q${b}"
        put(ws, f"A{r}", cfg)
        put(ws, f"B{r}", desc[cfg], wrap=True)
        if not dev:
            for col in "CDEFG":
                put(ws, f"{col}{r}", "–")
            continue
        put(ws, f"C{r}", f"=INDEX({E},MATCH(MIN({Rr}),{Rr},0))", fmt=R2)
        put(ws, f"D{r}", f"=MIN({Rr})", fmt=R3)
        put(ws, f"E{r}", f"=INDEX({E},MATCH(MAX({Qq}),{Qq},0))", fmt=R2)
        put(ws, f"F{r}", f"=MAX({Qq})", fmt=R3)
        put(ws, f"G{r}", f'=IF($B$8="F1-max",E{r},C{r})', fmt=R2)
    if not dev:
        put(ws, "A20", "Không có file dev (--dev-ratio 0): không chọn được ngưỡng, B11 dùng ngưỡng cố định B10.",
            F(color="C00000"))

    # B
    put(ws, "A22", "B. Tầng B trên test (cấu hình B7, ngưỡng B11)", sect)
    header(ws, 23, ["Chế độ", "Vùng bin", "TP", "FN", "FP", "TN", "MR", "FAR", "Precision", "Recall", "F1",
                    "Specificity", "DCF", "CI95 F1 (Python)", "CI95 DCF (Python)", "CI95 MR (Python)",
                    "CI95 FAR (Python)"])
    brows = [(24, "Strict (mọi bin)", "all", "$B$11"), (25, "Collar ±1 bin (bỏ bin biên)", "nb", "$B$11"),
             (26, "Chỉ bin biên", "b", "$B$11")]
    if dev:
        brows.append((27, "Tham khảo: ngưỡng chọn trên dev (B7, B8), mọi bin", "all",
                      "INDEX($G$17:$G$19,MATCH($B$7,$A$17:$A$19,0))"))
    brows.append((28, "Tham khảo: ngưỡng cố định B10, mọi bin", "all", "$B$10"))
    for r, lab, reg, thr_ref in brows:
        put(ws, f"A{r}", lab, wrap=True)
        put(ws, f"B{r}", reg)
        for col, src in zip("CDEF", "FGHI"):
            put(ws, f"{col}{r}", sumifs(src, f"$B{r}", thr_ref))
        put(ws, f"G{r}", f'=IF(C{r}+D{r}=0,"",D{r}/(C{r}+D{r}))', fmt=PCT)
        put(ws, f"H{r}", f'=IF(E{r}+F{r}=0,"",E{r}/(E{r}+F{r}))', fmt=PCT)
        put(ws, f"I{r}", f'=IF(C{r}+E{r}=0,"",C{r}/(C{r}+E{r}))', fmt=PCT)
        put(ws, f"J{r}", f'=IF(C{r}+D{r}=0,"",C{r}/(C{r}+D{r}))', fmt=PCT)
        put(ws, f"K{r}", f'=IF(2*C{r}+D{r}+E{r}=0,"",2*C{r}/(2*C{r}+D{r}+E{r}))', fmt=R3)
        put(ws, f"L{r}", f'=IF(E{r}+F{r}=0,"",F{r}/(E{r}+F{r}))', fmt=PCT)
        put(ws, f"M{r}", f'=IF(OR(G{r}="",H{r}=""),"",$B$6*G{r}+(1-$B$6)*H{r})', fmt=R3)
    fmt_ci = lambda t, pct=False: "" if not t else (  # noqa: E731
        f"{t[0]:.1%}–{t[1]:.1%}" if pct else f"{t[0]:.3f}–{t[1]:.3f}")
    if ci:
        put(ws, "N24", fmt_ci(ci["f1"]), blue)
        put(ws, "O24", fmt_ci(ci["dcf"]), blue)
        put(ws, "P24", fmt_ci(ci["mr"], True), blue)
        put(ws, "Q24", fmt_ci(ci["far"], True), blue)
        put(ws, "N25", fmt_ci(ci["f1_nb"]), blue)
        ws["N24"].comment = Comment(f"Bootstrap {args.bootstrap} lần theo file, seed {args.seed}, ở cấu hình "
                                    f"{sel['cfg']} / ngưỡng {sel['thr']}.", "excel_report.py")

    # C
    r0 = 30
    put(ws, f"A{r0 - 1}", "C. Theo category (test; cấu hình B7, ngưỡng B11)", sect)
    header(ws, r0, ["Category", "#file", "TP", "FN", "FP", "TN", "FEC", "MSC", "OVER", "NDS", "TP ±1", "FN ±1",
                    "FP ±1", "TN ±1", "% speech", "MR", "FAR", "Precision", "F1 strict", "F1 ±1 bin", "DCF",
                    "Lỗi biên (FEC+OVER)", "Lỗi thật (MSC+NDS)", "F1 CI95 thấp (Python)", "F1 CI95 cao (Python)"])

    def metric_cols(r):
        put(ws, f"O{r}", f'=IF(SUM(C{r}:F{r})=0,"",(C{r}+D{r})/SUM(C{r}:F{r}))', fmt=PCT)
        put(ws, f"P{r}", f'=IF(C{r}+D{r}=0,"",D{r}/(C{r}+D{r}))', fmt=PCT)
        put(ws, f"Q{r}", f'=IF(E{r}+F{r}=0,"",E{r}/(E{r}+F{r}))', fmt=PCT)
        put(ws, f"R{r}", f'=IF(C{r}+E{r}=0,"",C{r}/(C{r}+E{r}))', fmt=PCT)
        put(ws, f"S{r}", f'=IF(2*C{r}+D{r}+E{r}=0,"",2*C{r}/(2*C{r}+D{r}+E{r}))', fmt=R3)
        put(ws, f"T{r}", f'=IF(2*K{r}+L{r}+M{r}=0,"",2*K{r}/(2*K{r}+L{r}+M{r}))', fmt=R3)
        put(ws, f"U{r}", f'=IF(OR(P{r}="",Q{r}=""),"",$B$6*P{r}+(1-$B$6)*Q{r})', fmt=R3)
        put(ws, f"V{r}", f"=G{r}+I{r}")
        put(ws, f"W{r}", f"=H{r}+J{r}")

    cat_row = {}
    for i, cat in enumerate(cats):
        r = r0 + 1 + i
        cat_row[cat] = r
        put(ws, f"A{r}", cat)
        put(ws, f"B{r}", f"=COUNTIFS('DL theo file'!$B$5:$B${df_last},$A{r},'DL theo file'!$C$5:$C${df_last},\"test\")")
        for col, src in zip("CDEFGHIJ", "FGHIJKLM"):
            put(ws, f"{col}{r}", sumifs(src, '"all"', "$B$11", f"$A{r}"))
        for col, src in zip("KLMN", "FGHI"):
            put(ws, f"{col}{r}", sumifs(src, '"nb"', "$B$11", f"$A{r}"))
        metric_cols(r)
        if cat in ci_cat:
            put(ws, f"X{r}", ci_cat[cat][0], blue, R3)
            put(ws, f"Y{r}", ci_cat[cat][1], blue, R3)
    rt = r0 + 1 + len(cats)
    a, b = r0 + 1, rt - 1
    put(ws, f"A{rt}", "Tổng test (gộp)", F(bold=True))
    for col in "BCDEFGHIJKLMN":
        put(ws, f"{col}{rt}", f"=SUM({col}{a}:{col}{b})")
    metric_cols(rt)
    if ci:
        put(ws, f"X{rt}", ci["f1"][0], blue, R3)
        put(ws, f"Y{rt}", ci["f1"][1], blue, R3)
    rm = rt + 1
    put(ws, f"A{rm}", f"Macro (TB {len(cats)} category)", F(bold=True))
    for col, fm in zip("OPQRSTU", (PCT, PCT, PCT, PCT, R3, R3, R3)):
        put(ws, f"{col}{rm}", f"=AVERAGE({col}{a}:{col}{b})", fmt=fm)

    # D. YAMNet
    r = rm + 3
    if yam:
        put(ws, f"A{r}", f"D. So với YAMNet (sheet KQ du lieu that: cấu hình {yam['cfg']}, ngưỡng {yam['thr']}; "
                         f"số đếm nhập từ file, MR…DCF tính lại theo B6)", sect)
        r += 1
        header(ws, r, ["Category", "#file YAMNet", "#file FireRed", "TP (Y)", "FN (Y)", "FP (Y)", "TN (Y)",
                       "MR YAMNet", "MR FireRed", "FAR YAMNet", "FAR FireRed", "F1 YAMNet", "F1 FireRed",
                       "DCF YAMNet", "DCF FireRed", "Δ F1 (FR − Y)", "Δ FAR (FR − Y)", "Δ DCF (FR − Y)"])
        rows_d = [(c, yam["cats"][c], yam["nfile"].get(c), cat_row.get(c)) for c in yam["cats"]]
        rows_d.append(("Tổng test (gộp)", yam["strict"], sum(v or 0 for v in yam["nfile"].values()) or None, rt))
        for name, cnt, nfy, fr_row in rows_d:
            r += 1
            put(ws, f"A{r}", name, F(bold=name.startswith("Tổng")))
            put(ws, f"B{r}", nfy, blue)
            put(ws, f"C{r}", f"=B{fr_row}" if fr_row else None)
            for col, v in zip("DEFG", cnt):
                put(ws, f"{col}{r}", v, blue)
            put(ws, f"H{r}", f'=IF(D{r}+E{r}=0,"",E{r}/(D{r}+E{r}))', fmt=PCT)
            put(ws, f"J{r}", f'=IF(F{r}+G{r}=0,"",F{r}/(F{r}+G{r}))', fmt=PCT)
            put(ws, f"L{r}", f'=IF(2*D{r}+E{r}+F{r}=0,"",2*D{r}/(2*D{r}+E{r}+F{r}))', fmt=R3)
            put(ws, f"N{r}", f'=IF(OR(H{r}="",J{r}=""),"",$B$6*H{r}+(1-$B$6)*J{r})', fmt=R3)
            if fr_row:
                put(ws, f"I{r}", f"=P{fr_row}", fmt=PCT)
                put(ws, f"K{r}", f"=Q{fr_row}", fmt=PCT)
                put(ws, f"M{r}", f"=S{fr_row}", fmt=R3)
                put(ws, f"O{r}", f"=U{fr_row}", fmt=R3)
                put(ws, f"P{r}", f'=IF(OR(M{r}="",L{r}=""),"",M{r}-L{r})', fmt=R3)
                put(ws, f"Q{r}", f'=IF(OR(K{r}="",J{r}=""),"",K{r}-J{r})', fmt=PCT)
                put(ws, f"R{r}", f'=IF(OR(O{r}="",N{r}=""),"",O{r}-N{r})', fmt=R3)
        r += 1
        put(ws, f"A{r}", "So sánh chỉ công bằng khi #file hai cột bằng nhau (cùng tập test). Δ âm ở FAR / DCF "
                         "= FireRed tốt hơn; Δ dương ở F1 = FireRed tốt hơn.", note)
        r += 2
    # notes
    put(ws, f"A{r}", "Ghi chú", sect)
    notes = [
        f"File trong danh sách chia nhưng không chạy được: {len(problems)}"
        + (" – " + "; ".join(f"{n} ({why})" for n, why in problems[:15]) + (" …" if len(problems) > 15 else "")
           if problems else ""),
        "FEC = FN ở bin biên, MSC = FN xa biên, OVER = FP ở bin biên, NDS = FP xa biên.",
        "Bộ phân loại 'luôn speech' có DCF = 1 − B6 (= 0.25 với w 0.75): DCF test gần mức này nghĩa là ngưỡng "
        "đang gần như gọi mọi bin là speech.",
    ]
    for t in notes:
        r += 1
        put(ws, f"A{r}", t, note)

    widths = {"A": 30, "B": 16}
    for j in range(3, 26):
        widths[openpyxl.utils.get_column_letter(j)] = 11
    widths["N"] = widths["O"] = widths["P"] = widths["Q"] = 13
    for k, v in widths.items():
        ws.column_dimensions[k].width = v
    for rr in (16, 23, r0):
        ws.row_dimensions[rr].height = 42
    for rr in (17, 18, 19):
        ws.row_dimensions[rr].height = 40
    wb.calculation.fullCalcOnLoad = True           # Excel tính lại công thức khi mở
    wb.save(path)


# =========================================================================== main
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx", help="file tổng hợp YAMNet (lấy chia dev/test ở sheet 'DL theo file' + số YAMNet để so)")
    ap.add_argument("--split-csv", help="thay cho --xlsx: CSV có cột File, Category, Tập (dev/test)")
    ap.add_argument("--dev-ratio", type=float, default=0.3,
                    help="không có --xlsx / --split-csv: tỉ lệ file mỗi category đưa vào dev (0 = tất cả là test)")
    ap.add_argument("--split-seed", type=int, default=3)
    ap.add_argument("--threshold", type=float,
                    help="ngưỡng cố định cho mục B, C và CI (bỏ qua việc chọn trên dev). Không có dev mà không "
                         "nhập thì dùng 0.4")
    ap.add_argument("--audio-root", required=True)
    ap.add_argument("--gt-root", required=True, help="thư mục GT; ghép theo tên file, ưu tiên file trong ground_truth/")
    ap.add_argument("--quality", help="nếu một tên file có nhiều bản audio: chọn bản có chuỗi này trong đường dẫn")
    ap.add_argument("-o", "--out", default=os.path.join(HERE, "out", "firered_report.xlsx"))
    ap.add_argument("--cache", default=os.path.join(HERE, "out", "probs_cache"))
    ap.add_argument("--force", action="store_true", help="bỏ cache, chạy lại model")
    ap.add_argument("--model-dir", default=os.path.join(HERE, "models"))
    ap.add_argument("--onnx")
    ap.add_argument("--cmvn")
    ap.add_argument("--mode", default="auto", choices=["auto", "nonstream", "stream", "stream_cached"])
    ap.add_argument("--threads", type=int, default=1)
    ap.add_argument("--bin", type=float, default=0.5)
    ap.add_argument("--boundary", type=int, default=1, help="số bin mỗi phía điểm chuyển GT tính là biên")
    ap.add_argument("--rebin", choices=["overlap", "midpoint"], default="overlap")
    ap.add_argument("--min-overlap", type=float, default=0.2, help="--rebin overlap: tỉ lệ bin phải được phủ")
    ap.add_argument("--dcf-weight", type=float, default=0.75)
    ap.add_argument("--config", choices=CONFIGS, default="F1", help="cấu hình cho mục B, C và CI")
    ap.add_argument("--criterion", choices=["DCF-min", "F1-max"], default="DCF-min")
    ap.add_argument("--f2-pad-before", type=float, default=0.10)
    ap.add_argument("--f2-pad-after", type=float, default=0.12)
    ap.add_argument("--f2-drop", type=float, default=0.0)
    ap.add_argument("--f2-merge", type=float, default=0.5)
    ap.add_argument("--f2-order", default="pad,drop,merge")
    ap.add_argument("--bootstrap", type=int, default=1000)
    ap.add_argument("--seed", type=int, default=0)
    ap.add_argument("--limit", type=int, help="chỉ dùng N file đầu (thử nhanh)")
    args = ap.parse_args()
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(errors="replace")      # console Windows không có font tiếng Việt
        except Exception:
            pass
    args.f2_order = [s.strip() for s in args.f2_order.split(",") if s.strip()]
    if args.threshold is not None and not 0 < args.threshold < 1:
        raise SystemExit("--threshold phải nằm trong (0, 1)")
    args.fixed_thr = round(args.threshold, 4) if args.threshold is not None else 0.4
    if args.fixed_thr not in THRESHOLDS:          # ngưỡng lẻ (vd. 0.6026): thêm vào lưới quét
        THRESHOLDS.append(args.fixed_thr)
        THRESHOLDS.sort()

    if args.split_csv:
        split, args.split_source = read_split_csv(args.split_csv), os.path.basename(args.split_csv)
    elif args.xlsx:
        split, args.split_source = read_split_xlsx(args.xlsx), f"{os.path.basename(args.xlsx)} (sheet DL theo file)"
    else:
        split = auto_split(args.audio_root, args.gt_root, args.dev_ratio, args.split_seed)
        sp = os.path.splitext(os.path.abspath(args.out))[0] + "_split.csv"
        os.makedirs(os.path.dirname(sp), exist_ok=True)
        write_split(sp, split)
        args.split_source = (f"tự chia: {args.dev_ratio:.0%} dev theo file, phân tầng category, seed "
                             f"{args.split_seed} (lưu ở {os.path.basename(sp)})")
        print(f"Tự chia dev/test -> {sp} (dùng lại bằng --split-csv)")
    split = [s for s in split if s["split"] in ("dev", "test")]
    if args.limit:
        split = split[:args.limit]
    items, problems = match(split, args.audio_root, args.gt_root, args.quality)
    print(f"{len(items)}/{len(split)} file ghép được audio + GT"
          + (f"; {len(problems)} file lỗi, vd. {problems[:3]}" if problems else ""))
    if not items:
        raise SystemExit("Không có file nào để chạy")

    from firered_onnx import PostConfig

    post = PostConfig()
    get_probs(items, args)
    for it in items:
        if not it["category"]:
            it["category"] = gt_category(it["gt"], args.gt_root)
        evaluate_file(it, args, post)

    sel = {"cfg": args.config, "crit": args.criterion}
    k, thr = choose_threshold(items, args.config, args.criterion, args.dcf_weight)
    if args.threshold is not None or k is None:
        if k is None and args.threshold is None:
            print(f"[!] không có file dev: dùng ngưỡng cố định {args.fixed_thr} (mặc định FireRed)")
        k, thr = THRESHOLDS.index(args.fixed_thr), args.fixed_thr
        sel["mode"] = "cố định"
    else:
        sel["mode"] = "dev"
    sel.update(k=k, thr=thr)
    rng = np.random.default_rng(args.seed)
    test = [it for it in items if it["split"] == "test"]
    pf = np.array([it["counts"][args.config][k] for it in test]) if test else np.zeros((0, 2, 4))
    ci = bootstrap(pf, args.dcf_weight, args.bootstrap, rng) if len(pf) else {}
    ci_cat = {}
    for cat in sorted({it["category"] for it in test}):
        sub = np.array([it["counts"][args.config][k] for it in test if it["category"] == cat])
        bb = bootstrap(sub, args.dcf_weight, args.bootstrap, rng)
        if bb:
            ci_cat[cat] = bb["f1"]
    yam = read_yamnet(args.xlsx) if args.xlsx else None

    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    write_xlsx(args.out, items, problems, args, sel, ci, ci_cat, yam)

    tot = pf.sum(0).sum(0) if len(pf) else np.zeros(4)
    m = metr(*tot, args.dcf_weight)
    print(f"Ngưỡng cố định: {thr}" if sel["mode"] == "cố định"
          else f"Ngưỡng {args.criterion} trên dev ({args.config}): {thr}")
    print(f"Test strict: MR {m['mr']:.1%}  FAR {m['far']:.1%}  F1 {m['f1']:.3f}  DCF {m['dcf']:.3f}"
          + (f"  (CI F1 {ci['f1'][0]:.3f}–{ci['f1'][1]:.3f})" if ci else ""))
    print(f"-> {args.out}")


if __name__ == "__main__":
    main()
