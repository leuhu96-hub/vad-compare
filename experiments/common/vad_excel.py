"""Phần chung cho báo cáo Excel của các model VAD, cùng bố cục sheet "KQ du lieu that" (mục A, B, C)
trong file Tong_hop_danh_gia_VAD_test.xlsx của YAMNet. Mỗi model chỉ cần một "spec":

    class Spec:
        name, short, abbr, sheet          # tên hiển thị
        configs = ["X0", "X1", "X2"]       # đúng 3 cấu hình (bố cục Excel cố định)
        default_config, default_thr        # cấu hình cho mục B/C và ngưỡng mặc định của model
        frame_desc                         # vd. "10 ms" (dòng ghi chú A4)
        here                               # thư mục mặc định cho out/
        add_args(ap); prepare(args)        # tham số riêng của model
        model_label(args) -> str
        descriptions(args) -> {cfg: str};  pp_note(args) -> str
        load_probs(items, args)            # gán it["probs"], it["dur"] (+ khoá riêng), có cache
        decide(it, args, thr, n) -> {cfg: bool[n_thr, n]}   # quyết định theo bin GT cho mọi ngưỡng

Quy trình giống file YAMNet: chia dev/test (--xlsx / --split-csv / tự chia), GT bin 0.5 s, biên ±1 bin,
quét ngưỡng 0.01 … 0.95, chọn ngưỡng trên dev (DCF-min / F1-max) hoặc ngưỡng cố định (--threshold),
DCF = w·MR + (1−w)·FAR, CI bootstrap theo file.
"""
from __future__ import annotations

import argparse
import csv
import datetime
import os
import sys
from collections import defaultdict

import numpy as np

COMMON = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(COMMON, "..", "firered_vad"))      # dataset_reader.py nằm ở đây
from dataset_reader import AUDIO_EXTS, file_key, find_audio, gt_to_bins, read_ground_truth  # noqa: E402

THRESHOLDS = [round(k * 0.01, 2) for k in range(1, 96)]
REGIONS = ("all", "nb", "b")


# =========================================================================== split / matching
def read_split_xlsx(path: str):
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "DL theo file" not in wb.sheetnames:
        raise SystemExit(f"{path}: không có sheet 'DL theo file' để lấy danh sách dev/test")
    ws = wb["DL theo file"]
    rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True))
    hdr = next((i for i, r in enumerate(rows) if r and "File" in r and "Tập" in r), None)
    if hdr is None:
        raise SystemExit(f"{path}: sheet 'DL theo file' không có header 'File' / 'Category' / 'Tập'")
    h = rows[hdr]
    ci, cc, cs = h.index("File"), h.index("Category"), h.index("Tập")
    out = []
    for r in rows[hdr + 1:]:
        if not r or not r[ci]:
            continue
        out.append({"file": str(r[ci]).strip(), "category": str(r[cc]).strip(), "split": str(r[cs]).strip()})
    return out


def read_split_csv(path: str):
    out = []
    with open(path, encoding="utf-8-sig") as f:
        for r in csv.DictReader(f):
            low = {k.strip().lower(): (v or "").strip() for k, v in r.items()}
            name = low.get("file") or low.get("name")
            if not name:
                continue
            out.append({"file": name, "category": low.get("category", ""),
                        "split": low.get("tập") or low.get("split") or low.get("set") or "test"})
    return out


def index_gt(gt_root: str):
    """stem -> [paths]; nếu một tên có nhiều file, ưu tiên file nằm trong thư mục tên ground_truth."""
    idx = defaultdict(list)
    for dp, _, files in os.walk(gt_root):
        for f in files:
            if f.lower().endswith(".txt"):
                idx[file_key(f)].append(os.path.join(dp, f))
    out = {}
    for k, paths in idx.items():
        pref = [p for p in paths if "ground_truth" in [x.lower() for x in os.path.abspath(p).split(os.sep)]]
        out[k] = pref or paths
    return out


def match(split, audio_root, gt_root, quality=None, log=print):
    audio = defaultdict(list)
    for p, q in find_audio(audio_root, AUDIO_EXTS):
        audio[file_key(p)].append(p)
    gts = index_gt(gt_root)
    items, problems = [], []
    for s in split:
        k = file_key(s["file"])
        a = audio.get(k, [])
        if quality and len(a) > 1:
            a = [p for p in a if quality.lower() in p.lower()] or a
        g = gts.get(k, [])
        if not a or not g:
            problems.append((s["file"], "không có audio" if not a else "không có GT"))
            continue
        if len(g) > 1:
            problems.append((s["file"], f"{len(g)} file GT trùng tên: {g[:2]}"))
            continue
        if len(a) > 1:
            log(f"[!] {s['file']}: {len(a)} file audio trùng tên, dùng {a[0]} (chọn bản khác bằng --quality)")
        items.append({**s, "audio": a[0], "gt": g[0]})
    return items, problems


def gt_category(gt_path, gt_root):
    """category = thư mục đầu tiên dưới gt_root (Groundtruth/<category>/ground_truth/x.txt);
    GT để phẳng thì lấy tiền tố tên file trước dấu '_' (Asm_xxx -> Asm)."""
    rel = os.path.relpath(gt_path, gt_root).split(os.sep)
    if len(rel) > 1 and rel[0].lower() != "ground_truth":
        return rel[0]
    stem = os.path.splitext(os.path.basename(gt_path))[0]
    return stem.split("_", 1)[0] if "_" in stem else "Tất cả"


def auto_split(audio_root, gt_root, dev_ratio, seed):
    """Không có file chia sẵn: mọi file có cả audio và GT; chia dev/test theo file, phân tầng theo category."""
    audio = {file_key(p) for p, _ in find_audio(audio_root, AUDIO_EXTS)}
    by_cat = defaultdict(list)
    for k, paths in sorted(index_gt(gt_root).items()):
        if k in audio:
            name = os.path.splitext(os.path.basename(paths[0]))[0]
            by_cat[gt_category(paths[0], gt_root)].append(name)
    rng = np.random.default_rng(seed)
    out = []
    for cat in sorted(by_cat):
        names = sorted(by_cat[cat])
        n_dev = int(round(dev_ratio * len(names)))
        if dev_ratio > 0 and len(names) >= 2:
            n_dev = min(max(n_dev, 1), len(names) - 1)
        dev = set(rng.permutation(len(names))[:n_dev].tolist())
        out += [{"file": n, "category": cat, "split": "dev" if i in dev else "test"} for i, n in enumerate(names)]
    return out


def write_split(path, split):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["File", "Category", "Tập"])
        for s in split:
            w.writerow([s["file"], s["category"], s["split"]])


# =========================================================================== bins / counts
def boundary_mask(labels, tol=1):
    m = np.zeros(len(labels), bool)
    if tol <= 0:
        return m
    for t in range(1, len(labels)):
        a, b = labels[t - 1], labels[t]
        if a >= 0 and b >= 0 and a != b:
            m[max(0, t - tol):min(len(labels), t + tol)] = True
    return m


def counts(D, labels, bnd):
    """D [n_thr, n] bool -> [n_thr, 2 (b, nb), 4 (TP, FN, FP, TN)]."""
    y = labels == 1
    valid = labels >= 0
    out = np.zeros((D.shape[0], 2, 4), np.int64)
    for r, m in enumerate((valid & bnd, valid & ~bnd)):
        sp, ns = y & m, ~y & m
        out[:, r, 0] = (D & sp).sum(1)
        out[:, r, 1] = (~D & sp).sum(1)
        out[:, r, 2] = (D & ns).sum(1)
        out[:, r, 3] = (~D & ns).sum(1)
    return out


def evaluate_file(it, args, spec):
    gt = read_ground_truth(it["gt"])
    labels = np.array(gt_to_bins(gt, args.bin, it["dur"]), dtype=int)
    n = len(labels)
    it["labels"], it["bnd"] = labels, boundary_mask(labels, args.boundary)
    if n == 0:
        it["counts"] = {c: np.zeros((len(THRESHOLDS), 2, 4), np.int64) for c in spec.configs}
        return
    D = spec.decide(it, args, np.array(THRESHOLDS), n)
    it["counts"] = {c: counts(np.asarray(D[c], bool), labels, it["bnd"]) for c in spec.configs}


# =========================================================================== metrics / bootstrap
def metr(tp, fn, fp, tn, w):
    tp, fn, fp, tn = (np.asarray(x, float) for x in (tp, fn, fp, tn))
    with np.errstate(divide="ignore", invalid="ignore"):
        mr = np.where(tp + fn > 0, fn / (tp + fn), np.nan)
        far = np.where(fp + tn > 0, fp / (fp + tn), np.nan)
        f1 = np.where(2 * tp + fn + fp > 0, 2 * tp / (2 * tp + fn + fp), np.nan)
    return {"mr": mr, "far": far, "f1": f1, "dcf": w * mr + (1 - w) * far}


def choose_threshold(items, cfg, crit, w):
    dev = [it for it in items if it["split"] == "dev"]
    if not dev:
        return None, None
    c = sum(it["counts"][cfg] for it in dev).sum(1)        # [n_thr, 4] (all)
    m = metr(c[:, 0], c[:, 1], c[:, 2], c[:, 3], w)
    k = int(np.nanargmax(m["f1"])) if crit == "F1-max" else int(np.nanargmin(m["dcf"]))
    return k, THRESHOLDS[k]


def bootstrap(per_file, w, n_boot, rng):
    """per_file [nf, 2, 4] -> dict metric -> (lo, hi) cho all và nb."""
    nf = len(per_file)
    if nf == 0:
        return {}
    idx = rng.integers(0, nf, (n_boot, nf))
    s = per_file[idx].sum(1)                               # [B, 2, 4]
    a = s.sum(1)
    m = metr(a[:, 0], a[:, 1], a[:, 2], a[:, 3], w)
    nb = s[:, 1]
    m_nb = metr(nb[:, 0], nb[:, 1], nb[:, 2], nb[:, 3], w)
    ci = lambda x: (float(np.nanpercentile(x, 2.5)), float(np.nanpercentile(x, 97.5)))  # noqa: E731
    return {"f1": ci(m["f1"]), "dcf": ci(m["dcf"]), "mr": ci(m["mr"]), "far": ci(m["far"]),
            "f1_nb": ci(m_nb["f1"])}


# =========================================================================== YAMNet reference
def read_yamnet(path):
    """Số đếm test của YAMNet từ sheet 'KQ du lieu that' (giá trị đã lưu). None nếu không đọc được."""
    import openpyxl

    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        ws = wb["KQ du lieu that"]
    except Exception:
        return None
    rows = list(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 80), values_only=True))
    out = {"cfg": None, "thr": None, "strict": None, "cats": {}, "nfile": {}}
    for r in rows:
        if r and r[0] == "Cấu hình hậu xử lý cho tầng B, category":
            out["cfg"] = r[1]
        if r and r[0] == "Ngưỡng dùng (tự chọn trên dev theo B7, B8)":
            out["thr"] = r[1]
        if r and r[0] == "Strict (mọi bin)" and all(isinstance(x, (int, float)) for x in r[2:6]):
            out["strict"] = [int(x) for x in r[2:6]]
    try:
        hc = next(i for i, r in enumerate(rows) if r and r[0] == "Category" and r[2] == "TP")
    except StopIteration:
        return out if out["strict"] else None
    for r in rows[hc + 1:]:
        if not r or not r[0] or str(r[0]).startswith(("Tổng", "Macro")):
            break
        if all(isinstance(x, (int, float)) for x in r[2:6]):
            out["cats"][str(r[0])] = [int(x) for x in r[2:6]]
            out["nfile"][str(r[0])] = r[1]
    return out if out["strict"] else None


# =========================================================================== Excel
def write_xlsx(path, items, problems, args, sel, ci, ci_cat, yam, spec):
    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    NAVY, BLUE, YEL = "1F4E78", "0000FF", "FFF2CC"
    F = lambda **k: Font(name="Arial", size=k.pop("size", 10), **k)  # noqa: E731
    title, sect = F(size=14, bold=True, color=NAVY), F(size=11, bold=True, color=NAVY)
    hdr_font, hdr_fill = F(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=NAVY)
    yel = PatternFill("solid", fgColor=YEL)
    blue, note = F(color=BLUE), F(size=9, color="595959")
    PCT, R3, R2 = "0.0%", "0.000", "0.00"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = spec.sheet
    SHEET = spec.sheet
    assert len(spec.configs) == 3, "bố cục Excel cố định 3 cấu hình"
    dq = wb.create_sheet("DL quet nguong")
    df = wb.create_sheet("DL theo file")
    Q = "'DL quet nguong'"

    def put(sheet, ref, v, font=None, fmt=None, fill=None, wrap=False):
        c = sheet[ref]
        c.value = v
        c.font = font or F()
        if fmt:
            c.number_format = fmt
        if fill:
            c.fill = fill
        if wrap:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        return c

    def header(sheet, row, names, col=1):
        for j, h in enumerate(names):
            c = sheet.cell(row, col + j, h)
            c.font, c.fill = hdr_font, hdr_fill
            c.alignment = Alignment(wrap_text=True, vertical="center")

    # ------------------------------------------------------------- DL quet nguong
    put(dq, "A1", f"Số đếm theo ngưỡng – {spec.name}", sect)
    put(dq, "A2", "Tập: dev / test (theo sheet DL theo file). Vùng bin: all = mọi bin, nb = bỏ bin biên ±1, "
                  "b = chỉ bin biên. Chữ xanh = số đếm từ Python; MR…DCF (chỉ dev) là công thức.", note)
    cols = ["Tập", "Cấu hình", "Category", "Vùng bin", "Ngưỡng", "TP", "FN", "FP", "TN", "FEC", "MSC", "OVER",
            "NDS", "MR", "FAR", "Precision", "F1", "DCF"]
    header(dq, 4, cols)
    row = 5
    dev_rows = {}
    cats = sorted({it["category"] for it in items})

    def agg(sel_items, cfg):
        if not sel_items:
            return np.zeros((len(THRESHOLDS), 2, 4), np.int64)
        return sum(it["counts"][cfg] for it in sel_items)

    def write_row(split, cfg, cat, region, k, c):
        nonlocal row
        b, nb = c[k, 0], c[k, 1]
        if region == "all":
            v = b + nb
            extra = [b[1], nb[1], b[2], nb[2]]
        elif region == "nb":
            v, extra = nb, [0, nb[1], 0, nb[2]]
        else:
            v, extra = b, [b[1], 0, b[2], 0]
        vals = [split, cfg, cat, region, THRESHOLDS[k]] + [int(x) for x in v] + [int(x) for x in extra]
        for j, x in enumerate(vals):
            cell = dq.cell(row, j + 1, x)
            cell.font = blue if j >= 5 else F()
        if split == "dev":
            r = row
            dq.cell(r, 14, f'=IF(F{r}+G{r}=0,"",G{r}/(F{r}+G{r}))')
            dq.cell(r, 15, f'=IF(H{r}+I{r}=0,"",H{r}/(H{r}+I{r}))')
            dq.cell(r, 16, f'=IF(F{r}+H{r}=0,"",F{r}/(F{r}+H{r}))')
            dq.cell(r, 17, f'=IF(2*F{r}+G{r}+H{r}=0,"",2*F{r}/(2*F{r}+G{r}+H{r}))')
            dq.cell(r, 18, f"=IF(OR(N{r}=\"\",O{r}=\"\"),\"\",'{SHEET}'!$B$6*N{r}+(1-'{SHEET}'!$B$6)*O{r})")
            for j, fm in zip(range(14, 19), (PCT, PCT, PCT, R3, R3)):
                dq.cell(r, j).font = F()
                dq.cell(r, j).number_format = fm
        row += 1

    dev = [it for it in items if it["split"] == "dev"]
    test = [it for it in items if it["split"] == "test"]
    for cfg in spec.configs:
        c = agg(dev, cfg)
        dev_rows[cfg] = (row, row + len(THRESHOLDS) - 1)
        for k in range(len(THRESHOLDS)):
            write_row("dev", cfg, "Tất cả", "all", k, c)
    for cfg in spec.configs:
        for cat in cats:
            c = agg([it for it in test if it["category"] == cat], cfg)
            for region in REGIONS:
                for k in range(len(THRESHOLDS)):
                    write_row("test", cfg, cat, region, k, c)
    last = row - 1
    dq.freeze_panes = "A5"
    for j, wdt in enumerate([7, 9, 10, 9, 8] + [8] * 13, 1):
        dq.column_dimensions[openpyxl.utils.get_column_letter(j)].width = wdt

    def rng(col):
        return f"{Q}!${col}$5:${col}${last}"

    def sumifs(col, region_ref, thr_ref, cat_ref=None):
        s = (f'=SUMIFS({rng(col)},{rng("A")},"test",{rng("B")},$B$7,{rng("D")},{region_ref},'
             f'{rng("E")},{thr_ref}')
        return s + (f",{rng('C')},{cat_ref})" if cat_ref else ")")

    # ------------------------------------------------------------- DL theo file
    put(df, "A1", f"Số liệu theo file ({len(items)} file)", sect)
    put(df, "A2", f"Chữ xanh = số từ excel_report.py. Cột I–O tính ở cấu hình {sel['cfg']} và ngưỡng "
                  f"{sel['thr']} (lựa chọn của Python, xem '{SHEET}'!B12).", note)
    header(df, 4, ["File", "Category", "Tập", "Thời lượng (s)", "% speech (GT)", "#bin có nhãn", "#chuyển trạng thái",
                   "% bin biên", "TP", "FN", "FP", "TN", "MR", "FAR", "F1", "Audio", "GT"])
    for i, it in enumerate(items, 5):
        lab = it["labels"]
        valid = lab >= 0
        nlab = int(valid.sum())
        trans = int(np.sum((lab[1:] != lab[:-1]) & (lab[1:] >= 0) & (lab[:-1] >= 0))) if len(lab) > 1 else 0
        c = it["counts"][sel["cfg"]][sel["k"]].sum(0)
        m = metr(*c, args.dcf_weight)
        vals = [it["file"], it["category"], it["split"], round(it["dur"], 3),
                float((lab == 1).sum() / nlab) if nlab else None, nlab, trans,
                float((it["bnd"] & valid).sum() / nlab) if nlab else None,
                *[int(x) for x in c], *[None if np.isnan(m[x]) else float(m[x]) for x in ("mr", "far", "f1")],
                it["audio"], it["gt"]]
        for j, v in enumerate(vals, 1):
            cell = df.cell(i, j, v)
            cell.font = blue if 4 <= j <= 15 else F()
            if j in (5, 8, 13, 14):
                cell.number_format = PCT
            elif j == 15:
                cell.number_format = R3
    df_last = 4 + len(items)
    for j, wdt in enumerate([24, 10, 7, 11, 11, 9, 10, 9, 7, 7, 7, 7, 8, 8, 8, 50, 50], 1):
        df.column_dimensions[openpyxl.utils.get_column_letter(j)].width = wdt
    df.freeze_panes = "A5"

    # ------------------------------------------------------------- main sheet
    put(ws, "A1", f"Kết quả đánh giá {spec.name} – chỉ số tính bằng công thức", title)
    put(ws, "A2", f"Nguồn: excel_report.py ({datetime.date.today():%d/%m/%Y}). {len(items)} file "
                  f"(dev {len(dev)} / test {len(test)}); GT: {args.gt_root}; model: {spec.model_label(args)}. "
                  f"Chia dev/test lấy từ {args.split_source}.", note)
    put(ws, "A3", "Nền vàng = tham số bạn chỉnh (B6:B8, B10, B14) · chữ đen = công thức · chữ xanh = số nhập từ Python "
                  "(bootstrap; không tính lại trong Excel). Đổi tham số thì bảng A, B, C, D tự tính lại.", note)
    rule = (f"phủ ≥ {args.min_overlap:.0%} bin" if args.rebin == "overlap" else "tâm bin nằm trong đoạn speech")
    put(ws, "A4", f"Quyết định {spec.frame_desc} → bin {args.bin:g} s: {rule}. Biên = ±{args.boundary} bin quanh "
                  f"điểm chuyển GT. {spec.pp_note(args)}", note)
    put(ws, "A5", "Tham số", sect)
    params = [
        (6, "Trọng số miss trong DCF (FA = 1 − ô này)", args.dcf_weight, R2, True),
        (7, "Cấu hình hậu xử lý cho tầng B, category", sel["cfg"], None, True),
        (8, "Tiêu chí chọn ngưỡng trên dev", sel["crit"], None, True),
        (9, "Quy tắc rebin (cố định khi chạy script)", rule, None, False),
        (10, f"Ngưỡng cố định (--threshold; mặc định {spec.default_thr} của {spec.short})", args.fixed_thr, "0.00##", True),
    ]
    for r, lab, v, fmt, edit in params:
        put(ws, f"A{r}", lab, wrap=True)
        put(ws, f"B{r}", v, fmt=fmt, fill=yel if edit else None)
    dv1 = DataValidation(type="list", formula1='"' + ",".join(spec.configs) + '"', allow_blank=False)
    dv2 = DataValidation(type="list", formula1='"DCF-min,F1-max"', allow_blank=False)
    ws.add_data_validation(dv1)
    ws.add_data_validation(dv2)
    dv1.add("B7")
    dv2.add("B8")
    put(ws, "A11", "Ngưỡng dùng (B14 = dev: chọn trên dev theo B7, B8; B14 = cố định: B10)", wrap=True)
    put(ws, "B11", '=IF($B$14="cố định",$B$10,INDEX($G$17:$G$19,MATCH($B$7,$A$17:$A$19,0)))' if dev else "=$B$10",
        F(bold=True), "0.00##")
    put(ws, "A14", "Cách lấy ngưỡng (dev | cố định)", wrap=True)
    put(ws, "B14", sel["mode"], fill=yel)
    dv3 = DataValidation(type="list", formula1='"dev,cố định"', allow_blank=False)
    ws.add_data_validation(dv3)
    dv3.add("B14")
    put(ws, "A12", "Ngưỡng Python đã dùng cho CI và sheet DL theo file", wrap=True)
    put(ws, "B12", sel["thr"], blue, "0.00##")
    put(ws, "A13", "Kiểm tra")
    put(ws, "B13", f'=IF(AND(ABS(B11-B12)<0.0000001,B7="{sel["cfg"]}",ABS(B6-{args.dcf_weight})<0.0000001),"Khớp",'
                   f'"KHÁC: CI và DL theo file vẫn ở {sel["cfg"]} / ngưỡng {sel["thr"]} / w {args.dcf_weight}")')

    put(ws, "A15", "A. Chọn ngưỡng trên dev (qua toàn pipeline; DCF theo trọng số B6)", sect)
    header(ws, 16, ["Cấu hình", "Mô tả", "Ngưỡng DCF-min (dev)", "DCF min (dev)", "Ngưỡng F1-max (dev)",
                    "F1 max (dev)", "Ngưỡng dùng (theo B8)"])
    desc = spec.descriptions(args)
    for r, cfg in zip((17, 18, 19), spec.configs):
        a, b = dev_rows[cfg]
        E, Rr, Qq = f"{Q}!$E${a}:$E${b}", f"{Q}!$R${a}:$R${b}", f"{Q}!$Q${a}:$Q${b}"
        put(ws, f"A{r}", cfg)
        put(ws, f"B{r}", desc[cfg], wrap=True)
        if not dev:
            for col in "CDEFG":
                put(ws, f"{col}{r}", "–")
            continue
        put(ws, f"C{r}", f"=INDEX({E},MATCH(MIN({Rr}),{Rr},0))", fmt=R2)
        put(ws, f"D{r}", f"=MIN({Rr})", fmt=R3)
        put(ws, f"E{r}", f"=INDEX({E},MATCH(MAX({Qq}),{Qq},0))", fmt=R2)
        put(ws, f"F{r}", f"=MAX({Qq})", fmt=R3)
        put(ws, f"G{r}", f'=IF($B$8="F1-max",E{r},C{r})', fmt=R2)
    if not dev:
        put(ws, "A20", "Không có file dev (--dev-ratio 0): không chọn được ngưỡng, B11 dùng ngưỡng cố định B10.",
            F(color="C00000"))

    # B
    put(ws, "A22", "B. Tầng B trên test (cấu hình B7, ngưỡng B11)", sect)
    header(ws, 23, ["Chế độ", "Vùng bin", "TP", "FN", "FP", "TN", "MR", "FAR", "Precision", "Recall", "F1",
                    "Specificity", "DCF", "CI95 F1 (Python)", "CI95 DCF (Python)", "CI95 MR (Python)",
                    "CI95 FAR (Python)"])
    brows = [(24, "Strict (mọi bin)", "all", "$B$11"), (25, "Collar ±1 bin (bỏ bin biên)", "nb", "$B$11"),
             (26, "Chỉ bin biên", "b", "$B$11")]
    if dev:
        brows.append((27, "Tham khảo: ngưỡng chọn trên dev (B7, B8), mọi bin", "all",
                      "INDEX($G$17:$G$19,MATCH($B$7,$A$17:$A$19,0))"))
    brows.append((28, "Tham khảo: ngưỡng cố định B10, mọi bin", "all", "$B$10"))
    for r, lab, reg, thr_ref in brows:
        put(ws, f"A{r}", lab, wrap=True)
        put(ws, f"B{r}", reg)
        for col, src in zip("CDEF", "FGHI"):
            put(ws, f"{col}{r}", sumifs(src, f"$B{r}", thr_ref))
        put(ws, f"G{r}", f'=IF(C{r}+D{r}=0,"",D{r}/(C{r}+D{r}))', fmt=PCT)
        put(ws, f"H{r}", f'=IF(E{r}+F{r}=0,"",E{r}/(E{r}+F{r}))', fmt=PCT)
        put(ws, f"I{r}", f'=IF(C{r}+E{r}=0,"",C{r}/(C{r}+E{r}))', fmt=PCT)
        put(ws, f"J{r}", f'=IF(C{r}+D{r}=0,"",C{r}/(C{r}+D{r}))', fmt=PCT)
        put(ws, f"K{r}", f'=IF(2*C{r}+D{r}+E{r}=0,"",2*C{r}/(2*C{r}+D{r}+E{r}))', fmt=R3)
        put(ws, f"L{r}", f'=IF(E{r}+F{r}=0,"",F{r}/(E{r}+F{r}))', fmt=PCT)
        put(ws, f"M{r}", f'=IF(OR(G{r}="",H{r}=""),"",$B$6*G{r}+(1-$B$6)*H{r})', fmt=R3)
    fmt_ci = lambda t, pct=False: "" if not t else (  # noqa: E731
        f"{t[0]:.1%}–{t[1]:.1%}" if pct else f"{t[0]:.3f}–{t[1]:.3f}")
    if ci:
        put(ws, "N24", fmt_ci(ci["f1"]), blue)
        put(ws, "O24", fmt_ci(ci["dcf"]), blue)
        put(ws, "P24", fmt_ci(ci["mr"], True), blue)
        put(ws, "Q24", fmt_ci(ci["far"], True), blue)
        put(ws, "N25", fmt_ci(ci["f1_nb"]), blue)
        ws["N24"].comment = Comment(f"Bootstrap {args.bootstrap} lần theo file, seed {args.seed}, ở cấu hình "
                                    f"{sel['cfg']} / ngưỡng {sel['thr']}.", "excel_report.py")

    # C
    r0 = 30
    put(ws, f"A{r0 - 1}", "C. Theo category (test; cấu hình B7, ngưỡng B11)", sect)
    header(ws, r0, ["Category", "#file", "TP", "FN", "FP", "TN", "FEC", "MSC", "OVER", "NDS", "TP ±1", "FN ±1",
                    "FP ±1", "TN ±1", "% speech", "MR", "FAR", "Precision", "F1 strict", "F1 ±1 bin", "DCF",
                    "Lỗi biên (FEC+OVER)", "Lỗi thật (MSC+NDS)", "F1 CI95 thấp (Python)", "F1 CI95 cao (Python)"])

    def metric_cols(r):
        put(ws, f"O{r}", f'=IF(SUM(C{r}:F{r})=0,"",(C{r}+D{r})/SUM(C{r}:F{r}))', fmt=PCT)
        put(ws, f"P{r}", f'=IF(C{r}+D{r}=0,"",D{r}/(C{r}+D{r}))', fmt=PCT)
        put(ws, f"Q{r}", f'=IF(E{r}+F{r}=0,"",E{r}/(E{r}+F{r}))', fmt=PCT)
        put(ws, f"R{r}", f'=IF(C{r}+E{r}=0,"",C{r}/(C{r}+E{r}))', fmt=PCT)
        put(ws, f"S{r}", f'=IF(2*C{r}+D{r}+E{r}=0,"",2*C{r}/(2*C{r}+D{r}+E{r}))', fmt=R3)
        put(ws, f"T{r}", f'=IF(2*K{r}+L{r}+M{r}=0,"",2*K{r}/(2*K{r}+L{r}+M{r}))', fmt=R3)
        put(ws, f"U{r}", f'=IF(OR(P{r}="",Q{r}=""),"",$B$6*P{r}+(1-$B$6)*Q{r})', fmt=R3)
        put(ws, f"V{r}", f"=G{r}+I{r}")
        put(ws, f"W{r}", f"=H{r}+J{r}")

    cat_row = {}
    for i, cat in enumerate(cats):
        r = r0 + 1 + i
        cat_row[cat] = r
        put(ws, f"A{r}", cat)
        put(ws, f"B{r}", f"=COUNTIFS('DL theo file'!$B$5:$B${df_last},$A{r},'DL theo file'!$C$5:$C${df_last},\"test\")")
        for col, src in zip("CDEFGHIJ", "FGHIJKLM"):
            put(ws, f"{col}{r}", sumifs(src, '"all"', "$B$11", f"$A{r}"))
        for col, src in zip("KLMN", "FGHI"):
            put(ws, f"{col}{r}", sumifs(src, '"nb"', "$B$11", f"$A{r}"))
        metric_cols(r)
        if cat in ci_cat:
            put(ws, f"X{r}", ci_cat[cat][0], blue, R3)
            put(ws, f"Y{r}", ci_cat[cat][1], blue, R3)
    rt = r0 + 1 + len(cats)
    a, b = r0 + 1, rt - 1
    put(ws, f"A{rt}", "Tổng test (gộp)", F(bold=True))
    for col in "BCDEFGHIJKLMN":
        put(ws, f"{col}{rt}", f"=SUM({col}{a}:{col}{b})")
    metric_cols(rt)
    if ci:
        put(ws, f"X{rt}", ci["f1"][0], blue, R3)
        put(ws, f"Y{rt}", ci["f1"][1], blue, R3)
    rm = rt + 1
    put(ws, f"A{rm}", f"Macro (TB {len(cats)} category)", F(bold=True))
    for col, fm in zip("OPQRSTU", (PCT, PCT, PCT, PCT, R3, R3, R3)):
        put(ws, f"{col}{rm}", f"=AVERAGE({col}{a}:{col}{b})", fmt=fm)

    # D. YAMNet
    r = rm + 3
    if yam:
        put(ws, f"A{r}", f"D. So với YAMNet (sheet KQ du lieu that: cấu hình {yam['cfg']}, ngưỡng {yam['thr']}; "
                         f"số đếm nhập từ file, MR…DCF tính lại theo B6)", sect)
        r += 1
        header(ws, r, ["Category", "#file YAMNet", f"#file {spec.short}", "TP (Y)", "FN (Y)", "FP (Y)", "TN (Y)",
                       "MR YAMNet", f"MR {spec.short}", "FAR YAMNet", f"FAR {spec.short}", "F1 YAMNet",
                       f"F1 {spec.short}", "DCF YAMNet", f"DCF {spec.short}", f"Δ F1 ({spec.abbr} − Y)",
                       f"Δ FAR ({spec.abbr} − Y)", f"Δ DCF ({spec.abbr} − Y)"])
        rows_d = [(c, yam["cats"][c], yam["nfile"].get(c), cat_row.get(c)) for c in yam["cats"]]
        rows_d.append(("Tổng test (gộp)", yam["strict"], sum(v or 0 for v in yam["nfile"].values()) or None, rt))
        for name, cnt, nfy, fr_row in rows_d:
            r += 1
            put(ws, f"A{r}", name, F(bold=name.startswith("Tổng")))
            put(ws, f"B{r}", nfy, blue)
            put(ws, f"C{r}", f"=B{fr_row}" if fr_row else None)
            for col, v in zip("DEFG", cnt):
                put(ws, f"{col}{r}", v, blue)
            put(ws, f"H{r}", f'=IF(D{r}+E{r}=0,"",E{r}/(D{r}+E{r}))', fmt=PCT)
            put(ws, f"J{r}", f'=IF(F{r}+G{r}=0,"",F{r}/(F{r}+G{r}))', fmt=PCT)
            put(ws, f"L{r}", f'=IF(2*D{r}+E{r}+F{r}=0,"",2*D{r}/(2*D{r}+E{r}+F{r}))', fmt=R3)
            put(ws, f"N{r}", f'=IF(OR(H{r}="",J{r}=""),"",$B$6*H{r}+(1-$B$6)*J{r})', fmt=R3)
            if fr_row:
                put(ws, f"I{r}", f"=P{fr_row}", fmt=PCT)
                put(ws, f"K{r}", f"=Q{fr_row}", fmt=PCT)
                put(ws, f"M{r}", f"=S{fr_row}", fmt=R3)
                put(ws, f"O{r}", f"=U{fr_row}", fmt=R3)
                put(ws, f"P{r}", f'=IF(OR(M{r}="",L{r}=""),"",M{r}-L{r})', fmt=R3)
                put(ws, f"Q{r}", f'=IF(OR(K{r}="",J{r}=""),"",K{r}-J{r})', fmt=PCT)
                put(ws, f"R{r}", f'=IF(OR(O{r}="",N{r}=""),"",O{r}-N{r})', fmt=R3)
        r += 1
        put(ws, f"A{r}", "So sánh chỉ công bằng khi #file hai cột bằng nhau (cùng tập test). Δ âm ở FAR / DCF "
                         f"= {spec.short} tốt hơn; Δ dương ở F1 = {spec.short} tốt hơn.", note)
        r += 2
    # notes
    put(ws, f"A{r}", "Ghi chú", sect)
    notes = [
        f"File trong danh sách chia nhưng không chạy được: {len(problems)}"
        + (" – " + "; ".join(f"{n} ({why})" for n, why in problems[:15]) + (" …" if len(problems) > 15 else "")
           if problems else ""),
        "FEC = FN ở bin biên, MSC = FN xa biên, OVER = FP ở bin biên, NDS = FP xa biên.",
        "Bộ phân loại 'luôn speech' có DCF = 1 − B6 (= 0.25 với w 0.75): DCF test gần mức này nghĩa là ngưỡng "
        "đang gần như gọi mọi bin là speech.",
    ]
    for t in notes:
        r += 1
        put(ws, f"A{r}", t, note)

    widths = {"A": 30, "B": 16}
    for j in range(3, 26):
        widths[openpyxl.utils.get_column_letter(j)] = 11
    widths["N"] = widths["O"] = widths["P"] = widths["Q"] = 13
    for k, v in widths.items():
        ws.column_dimensions[k].width = v
    for rr in (16, 23, r0):
        ws.row_dimensions[rr].height = 42
    for rr in (17, 18, 19):
        ws.row_dimensions[rr].height = 40
    wb.calculation.fullCalcOnLoad = True           # Excel tính lại công thức khi mở
    wb.save(path)


# =========================================================================== main
def main(spec, argv=None):
    ap = argparse.ArgumentParser(description=spec.doc, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx", help="file tổng hợp YAMNet (lấy chia dev/test ở sheet 'DL theo file' + số YAMNet để so)")
    ap.add_argument("--split-csv", help="thay cho --xlsx: CSV có cột File, Category, Tập (dev/test)")
    ap.add_argument("--dev-ratio", type=float, default=0.3,
                    help="không có --xlsx / --split-csv: tỉ lệ file mỗi category đưa vào dev (0 = tất cả là test)")
    ap.add_argument("--split-seed", type=int, default=3)
    ap.add_argument("--threshold", type=float,
                    help="ngưỡng cố định cho mục B, C và CI (bỏ qua việc chọn trên dev). Không có dev mà không "
                         f"nhập thì dùng {spec.default_thr}")
    ap.add_argument("--audio-root", required=True)
    ap.add_argument("--gt-root", required=True, help="thư mục GT; ghép theo tên file, ưu tiên file trong ground_truth/")
    ap.add_argument("--quality", help="nếu một tên file có nhiều bản audio: chọn bản có chuỗi này trong đường dẫn")
    ap.add_argument("-o", "--out", default=os.path.join(spec.here, "out", spec.out_name))
    ap.add_argument("--cache", default=os.path.join(spec.here, "out", "probs_cache"))
    ap.add_argument("--force", action="store_true", help="bỏ cache, chạy lại model")
    ap.add_argument("--threads", type=int, default=1)
    ap.add_argument("--bin", type=float, default=0.5)
    ap.add_argument("--boundary", type=int, default=1, help="số bin mỗi phía điểm chuyển GT tính là biên")
    ap.add_argument("--rebin", choices=["overlap", "midpoint"], default="overlap")
    ap.add_argument("--min-overlap", type=float, default=0.2, help="--rebin overlap: tỉ lệ bin phải được phủ")
    ap.add_argument("--dcf-weight", type=float, default=0.75)
    ap.add_argument("--config", choices=spec.configs, default=spec.default_config,
                    help="cấu hình cho mục B, C và CI")
    ap.add_argument("--criterion", choices=["DCF-min", "F1-max"], default="DCF-min")
    ap.add_argument("--bootstrap", type=int, default=1000)
    ap.add_argument("--seed", type=int, default=0)
    ap.add_argument("--limit", type=int, help="chỉ dùng N file đầu (thử nhanh)")
    spec.add_args(ap)
    args = ap.parse_args(argv)
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(errors="replace")      # console Windows không có font tiếng Việt
        except Exception:
            pass
    spec.prepare(args)
    if args.threshold is not None and not 0 < args.threshold < 1:
        raise SystemExit("--threshold phải nằm trong (0, 1)")
    args.fixed_thr = round(args.threshold, 4) if args.threshold is not None else spec.default_thr
    if args.fixed_thr not in THRESHOLDS:          # ngưỡng lẻ (vd. 0.6026): thêm vào lưới quét
        THRESHOLDS.append(args.fixed_thr)
        THRESHOLDS.sort()

    if args.split_csv:
        split, args.split_source = read_split_csv(args.split_csv), os.path.basename(args.split_csv)
    elif args.xlsx:
        split, args.split_source = read_split_xlsx(args.xlsx), f"{os.path.basename(args.xlsx)} (sheet DL theo file)"
    else:
        split = auto_split(args.audio_root, args.gt_root, args.dev_ratio, args.split_seed)
        sp = os.path.splitext(os.path.abspath(args.out))[0] + "_split.csv"
        os.makedirs(os.path.dirname(sp), exist_ok=True)
        write_split(sp, split)
        args.split_source = (f"tự chia: {args.dev_ratio:.0%} dev theo file, phân tầng category, seed "
                             f"{args.split_seed} (lưu ở {os.path.basename(sp)})")
        print(f"Tự chia dev/test -> {sp} (dùng lại bằng --split-csv)")
    split = [s for s in split if s["split"] in ("dev", "test")]
    if args.limit:
        split = split[:args.limit]
    items, problems = match(split, args.audio_root, args.gt_root, args.quality)
    print(f"{len(items)}/{len(split)} file ghép được audio + GT"
          + (f"; {len(problems)} file lỗi, vd. {problems[:3]}" if problems else ""))
    if not items:
        raise SystemExit("Không có file nào để chạy")

    spec.load_probs(items, args)
    for it in items:
        if not it["category"]:
            it["category"] = gt_category(it["gt"], args.gt_root)
        evaluate_file(it, args, spec)

    sel = {"cfg": args.config, "crit": args.criterion}
    k, thr = choose_threshold(items, args.config, args.criterion, args.dcf_weight)
    if args.threshold is not None or k is None:
        if k is None and args.threshold is None:
            print(f"[!] không có file dev: dùng ngưỡng cố định {args.fixed_thr} (mặc định {spec.short})")
        k, thr = THRESHOLDS.index(args.fixed_thr), args.fixed_thr
        sel["mode"] = "cố định"
    else:
        sel["mode"] = "dev"
    sel.update(k=k, thr=thr)
    rng = np.random.default_rng(args.seed)
    test = [it for it in items if it["split"] == "test"]
    pf = np.array([it["counts"][args.config][k] for it in test]) if test else np.zeros((0, 2, 4))
    ci = bootstrap(pf, args.dcf_weight, args.bootstrap, rng) if len(pf) else {}
    ci_cat = {}
    for cat in sorted({it["category"] for it in test}):
        sub = np.array([it["counts"][args.config][k] for it in test if it["category"] == cat])
        bb = bootstrap(sub, args.dcf_weight, args.bootstrap, rng)
        if bb:
            ci_cat[cat] = bb["f1"]
    yam = read_yamnet(args.xlsx) if args.xlsx else None

    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    write_xlsx(args.out, items, problems, args, sel, ci, ci_cat, yam, spec)

    tot = pf.sum(0).sum(0) if len(pf) else np.zeros(4)
    m = metr(*tot, args.dcf_weight)
    print(f"Ngưỡng cố định: {thr}" if sel["mode"] == "cố định"
          else f"Ngưỡng {args.criterion} trên dev ({args.config}): {thr}")
    print(f"Test strict: MR {m['mr']:.1%}  FAR {m['far']:.1%}  F1 {m['f1']:.3f}  DCF {m['dcf']:.3f}"
          + (f"  (CI F1 {ci['f1'][0]:.3f}–{ci['f1'][1]:.3f})" if ci else ""))
    print(f"-> {args.out}")
    return items, sel


# =========================================================================== tiện ích cho spec
def seg_bins(segs, n, dur, bin_s, rule, min_ov):
    """Đoạn speech [(start_s, end_s)] (đã sắp, không chồng) -> bool[n] theo bin GT."""
    edges = np.arange(n + 1) * bin_s
    hi = np.minimum(edges[1:], dur)
    if rule == "midpoint":
        mid = (edges[:-1] + hi) / 2
        if not len(segs):
            return np.zeros(n, bool)
        s = np.array(segs, float)
        k = np.searchsorted(s[:, 0], mid, side="right") - 1
        return (k >= 0) & (mid < s[np.maximum(k, 0), 1])
    if not len(segs):
        return np.zeros(n, bool)
    s = np.array(segs, float)
    L = s[:, 1] - s[:, 0]
    C = np.clip(edges[:, None] - s[None, :, 0], 0, L[None, :]).sum(1)     # thời lượng speech tới mỗi mốc
    cov = np.diff(C)
    return cov >= min_ov * (hi - edges[:-1]) - 1e-9


def merge_segs(segs, gap=0.0):
    """Hợp các đoạn; nối đoạn cách nhau < gap giây."""
    out = []
    for a, b in sorted(segs):
        if out and (a <= out[-1][1] or a - out[-1][1] < gap):
            out[-1][1] = max(out[-1][1], b)
        else:
            out.append([a, b])
    return [(a, b) for a, b in out]


def rules_segs(segs, dur, pad_b, pad_a, drop, merge, order):
    """Hậu xử lý kiểu YAMNet A4 trên đoạn (giây)."""
    segs = list(segs)
    for step in order:
        if not segs:
            break
        if step == "pad":
            segs = merge_segs([(max(0.0, a - pad_b), min(dur, b + pad_a)) for a, b in segs])
        elif step == "drop":
            segs = [(a, b) for a, b in segs if b - a >= drop - 1e-9]
        elif step == "merge":
            segs = merge_segs(segs, merge)
        else:
            raise ValueError(f"bước hậu xử lý lạ: {step}")
    return segs
